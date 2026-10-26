import base64, io, json, logging, zipfile
from datetime import timedelta
from contextlib import asynccontextmanager
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from sqlalchemy import select, delete, text
from sqlalchemy.exc import IntegrityError
from argon2.exceptions import VerifyMismatchError, VerificationError, InvalidHashError
from pypdf import PdfReader
from docx import Document
from .config import settings
from .db import *
from .schemas import ProfileData, Credentials, ApplicationData, Preferences
from .security import get_db, current_user, new_session, hasher, digest, encrypt, decrypt, rate_limit, redis, cipher
from .services import profile_for, evaluate_job, match_job
from .ai import extract_resume

app=FastAPI(title='InternMatch AI',version='0.1.0')
app.add_middleware(CORSMiddleware,allow_origins=[settings.frontend_origin],allow_credentials=True,allow_methods=['GET','POST','PUT','DELETE'],allow_headers=['Content-Type'])

@app.middleware('http')
async def guard(request,call_next):
    if request.method in ('POST','PUT','DELETE','PATCH') and request.headers.get('origin') != settings.frontend_origin:
        return JSONResponse({'detail':'A matching Origin header is required'},status_code=403)
    try:
        # Direct client address only; forwarded headers must be set by a trusted proxy.
        rate_limit('http:'+str(request.client.host if request.client else 'unknown'),limit=180)
    except HTTPException as e: return JSONResponse({'detail':e.detail},status_code=e.status_code)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['Cache-Control']='no-store'
    return response

@app.exception_handler(Exception)
async def unexpected(request,exc):
    logging.getLogger('internmatch').error('Request failed: %s',type(exc).__name__)
    return JSONResponse({'detail':'Operation unavailable; consult server logs and configuration'},status_code=503)

@app.get('/health')
def health(db=Depends(get_db)):
    db.execute(text('SELECT 1')); redis.ping(); cipher()
    return {'status':'ok','integrations':{'ai':bool(settings.ai_api_key),'discovery_sources':bool(settings.greenhouse_boards or settings.lever_sites),'email':bool(settings.smtp_host)}}

@app.post('/auth/register')
def register(body:Credentials,response:Response,request:Request,db=Depends(get_db)):
    rate_limit('register:'+request.client.host,5,3600)
    cipher()
    user=User(email=body.email,password=hasher.hash(body.password)); db.add(user)
    try: db.flush()
    except IntegrityError: db.rollback(); raise HTTPException(409,'Account cannot be created with these details')
    new_session(db,user,response); db.add(Audit(user_id=user.id,action='account.created')); db.commit()
    return {'id':user.id,'email':user.email}

@app.post('/auth/login')
def login(body:Credentials,response:Response,request:Request,db=Depends(get_db)):
    rate_limit('login:'+request.client.host,10,900)
    user=db.scalar(select(User).where(User.email==body.email))
    try:
        # Hash on missing users too, reducing account enumeration timing differences.
        if not user:
            hasher.verify(hasher.hash('unavailable-account-password'),body.password)
            raise VerifyMismatchError()
        hasher.verify(user.password,body.password)
    except (VerifyMismatchError,VerificationError,InvalidHashError): raise HTTPException(401,'Invalid credentials')
    new_session(db,user,response); db.add(Audit(user_id=user.id,action='account.login')); db.commit()
    return {'id':user.id,'email':user.email}

@app.get('/auth/me')
def me(user=Depends(current_user)): return {'id':user.id,'email':user.email}

@app.post('/auth/logout')
def logout(request:Request,response:Response,db=Depends(get_db)):
    token=request.cookies.get('internmatch_session','')
    db.execute(delete(AuthSession).where(AuthSession.token_hash==digest(token))); db.commit()
    response.delete_cookie('internmatch_session')
    return {'ok':True}

@app.get('/profile')
def get_profile(user=Depends(current_user),db=Depends(get_db)):
    _,p=profile_for(db,user.id); return p

@app.put('/profile')
def set_profile(body:ProfileData,user=Depends(current_user),db=Depends(get_db)):
    row=db.get(Profile,user.id)
    if row: row.encrypted_data=encrypt(body.model_dump(mode='json')); row.version+=1
    else: db.add(Profile(user_id=user.id,encrypted_data=encrypt(body.model_dump(mode='json'))))
    db.add(Audit(user_id=user.id,action='profile.confirmed')); db.commit()
    redis.set('refresh:'+user.id,'1',ex=86400)
    return body

@app.post('/resumes')
def upload_resume(file:UploadFile=File(...),user=Depends(current_user),db=Depends(get_db)):
    rate_limit('resume:'+user.id,10,3600)
    raw=file.file.read(5_000_001)
    if len(raw)>5_000_000: raise HTTPException(413,'Resume exceeds 5 MB')
    name=(file.filename or '').split('/')[-1].split('\\')[-1][:255]
    try:
        if name.lower().endswith('.pdf') and raw.startswith(b'%PDF-'):
            doc=PdfReader(io.BytesIO(raw))
            if doc.is_encrypted or len(doc.pages)>20: raise ValueError('Encrypted or oversized PDF')
            content='\n'.join(page.extract_text() or '' for page in doc.pages)
        elif name.lower().endswith('.docx') and raw.startswith(b'PK'):
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist())>20_000_000 or len(z.infolist())>1000: raise ValueError('Expanded document too large')
                if 'word/document.xml' not in z.namelist(): raise ValueError('Invalid DOCX')
            doc=Document(io.BytesIO(raw))
            content='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
        else: raise ValueError('Only PDF and DOCX files are accepted')
        if not content.strip() or len(content)>100000: raise ValueError('Document needs accessible text within 100,000 characters; OCR is not available')
    except Exception as exc: raise HTTPException(422,'Cannot parse resume: '+str(exc)[:150])
    extraction=extract_resume(content)
    row=Resume(user_id=user.id,filename=name,encrypted_document=encrypt({'base64':base64.b64encode(raw).decode()}),encrypted_extraction=encrypt(extraction.model_dump(mode='json')))
    db.add(row); db.add(Audit(user_id=user.id,action='resume.uploaded')); db.commit()
    return {'id':row.id,'filename':name,'confirmed':False,'extraction':extraction}

@app.get('/resumes')
def resumes(user=Depends(current_user),db=Depends(get_db)):
    return [{'id':x.id,'filename':x.filename,'confirmed':x.confirmed,'created_at':x.created_at,'extraction':decrypt(x.encrypted_extraction)} for x in db.scalars(select(Resume).where(Resume.user_id==user.id).order_by(Resume.created_at.desc()))]

def owned_resume(db,id,user_id):
    row=db.get(Resume,id)
    if not row or row.user_id!=user_id: raise HTTPException(404,'Resume not found')
    return row

@app.get('/resumes/{id}/download')
def download_resume(id:str,user=Depends(current_user),db=Depends(get_db)):
    row=owned_resume(db,id,user.id)
    db.add(Audit(user_id=user.id,action='resume.downloaded')); db.commit()
    return Response(base64.b64decode(decrypt(row.encrypted_document)['base64']),media_type='application/octet-stream',headers={'Content-Disposition':'attachment; filename="resume'+('.pdf' if row.filename.lower().endswith('.pdf') else '.docx')+'"'})

@app.post('/resumes/{id}/confirm')
def confirm_resume(id:str,body:ProfileData,user=Depends(current_user),db=Depends(get_db)):
    row=owned_resume(db,id,user.id)
    row.confirmed=True; row.encrypted_extraction=encrypt(body.model_dump(mode='json'))
    db.add(Qualification(user_id=user.id,resume_id=id,encrypted_data=encrypt(body.model_dump(mode='json'))))
    return set_profile(body,user,db)

@app.delete('/resumes/{id}')
def remove_resume(id:str,user=Depends(current_user),db=Depends(get_db)):
    row=owned_resume(db,id,user.id); db.delete(row); db.commit(); return {'ok':True}

@app.get('/jobs')
def jobs(mode:str='strict',q:str='',company:str='',location:str='',arrangement:str='',skill:str='',season:str='',sort:str='recent',offset:int=0,limit:int=30,user=Depends(current_user),db=Depends(get_db)):
    if mode not in ('strict','potential','explore'): raise HTTPException(422,'Invalid mode')
    if sort not in ('recent','company','fit','deadline'): raise HTTPException(422,'Invalid sort')
    if offset<0 or limit<1 or limit>100: raise HTTPException(422,'Invalid pagination')
    query=select(Job).where(Job.status=='open')
    for column,value in [(Job.title,q),(Job.company,company),(Job.location,location)]:
        if value: query=query.where(column.ilike('%'+value.replace('%','\\%').replace('_','\\_')+'%',escape='\\'))
    # Text matching is a discovery filter, never eligibility evidence.
    for value in (skill,season,arrangement):
        if value: query=query.where(Job.text.ilike('%'+value.replace('%','\\%').replace('_','\\_')+'%',escape='\\'))
    query=query.order_by(Job.company if sort=='company' else Job.first_seen_at.desc())
    hidden={x.job_id for x in db.scalars(select(SavedJob).where(SavedJob.user_id==user.id,SavedJob.hidden==True))}
    saved={x.job_id for x in db.scalars(select(SavedJob).where(SavedJob.user_id==user.id,SavedJob.hidden==False))}
    counts={x:0 for x in ['Verified Eligible','Potentially Eligible','Ineligible','Unverified']}; results=[]
    # Counts cover the filtered candidate set. Bounded catalogue scan reported explicitly.
    candidates=list(db.scalars(query.limit(1001))); truncated=len(candidates)>1000
    for job in candidates[:1000]:
        if job.id in hidden: continue
        result=evaluate_job(db,user.id,job); counts[result['status']]+=1
        if mode=='strict' and result['status']!='Verified Eligible': continue
        if mode=='potential' and result['status'] not in ('Verified Eligible','Potentially Eligible'): continue
        profile=db.get(Profile,user.id); version=profile.version if profile else 0
        match=db.scalar(select(MatchScore).where(MatchScore.user_id==user.id,MatchScore.job_id==job.id,MatchScore.profile_version==version,MatchScore.content_hash==job.content_hash))
        results.append({'id':job.id,'company':job.company,'title':job.title,'location':job.location,'url':job.url,'details':job.details,'retrieved_at':job.retrieved_at,'first_seen_at':job.first_seen_at,'eligibility':result,'fit':decrypt(match.encrypted_data) if match else None,'saved':job.id in saved})
    if sort=='fit': results.sort(key=lambda x: (x['fit'] or {}).get('score') or -1,reverse=True)
    if sort=='deadline': results.sort(key=lambda x:x['details'].get('deadline') or '9999')
    db.commit()
    return {'items':results[offset:offset+limit],'counts':counts,'total':len(results),'truncated':truncated,'offset':offset,'limit':limit}

def get_job(db,id):
    job=db.get(Job,id)
    if not job: raise HTTPException(404,'Posting not found')
    return job

@app.get('/jobs/{id}')
def job_detail(id:str,user=Depends(current_user),db=Depends(get_db)):
    job=get_job(db,id); result=evaluate_job(db,user.id,job); db.commit()
    return {'id':job.id,'title':job.title,'company':job.company,'url':job.url,'text':job.text,'content_hash':job.content_hash,'status':job.status,'retrieved_at':job.retrieved_at,'details':job.details,'eligibility':result,'extraction_error':job.extraction_error}

@app.post('/jobs/{id}/match')
def match(id:str,user=Depends(current_user),db=Depends(get_db)):
    rate_limit('match:'+user.id,20,3600)
    job=get_job(db,id); result=match_job(db,user.id,job); db.commit(); return result

@app.put('/jobs/{id}/saved')
def save(id:str,hidden:bool=False,user=Depends(current_user),db=Depends(get_db)):
    get_job(db,id); row=db.get(SavedJob,(user.id,id))
    if row: row.hidden=hidden
    else: db.add(SavedJob(user_id=user.id,job_id=id,hidden=hidden))
    db.commit(); return {'ok':True}

@app.delete('/jobs/{id}/saved')
def unsave(id:str,user=Depends(current_user),db=Depends(get_db)):
    db.execute(delete(SavedJob).where(SavedJob.user_id==user.id,SavedJob.job_id==id)); db.commit(); return {'ok':True}

@app.get('/applications')
def applications(user=Depends(current_user),db=Depends(get_db)):
    return [{'id':x.id,'job_id':x.job_id,'company':db.get(Job,x.job_id).company,'title':db.get(Job,x.job_id).title,'url':db.get(Job,x.job_id).url,**decrypt(x.encrypted_data),'stage':x.stage} for x in db.scalars(select(Application).where(Application.user_id==user.id))]

@app.put('/applications/{job_id}')
def application(job_id:str,body:ApplicationData,user=Depends(current_user),db=Depends(get_db)):
    get_job(db,job_id)
    if body.resume_id: owned_resume(db,body.resume_id,user.id)
    row=db.scalar(select(Application).where(Application.user_id==user.id,Application.job_id==job_id))
    if row: row.stage=body.stage; row.encrypted_data=encrypt(body.model_dump(mode='json'))
    else: db.add(Application(user_id=user.id,job_id=job_id,stage=body.stage,encrypted_data=encrypt(body.model_dump(mode='json'))))
    db.commit(); return {'ok':True}

@app.get('/notifications')
def notifications(user=Depends(current_user),db=Depends(get_db)):
    return [{'id':x.id,'data':x.data,'read':x.read,'created_at':x.created_at} for x in db.scalars(select(Notification).where(Notification.user_id==user.id).order_by(Notification.created_at.desc()).limit(100))]

@app.put('/notifications/{id}/read')
def mark_read(id:str,user=Depends(current_user),db=Depends(get_db)):
    row=db.get(Notification,id)
    if not row or row.user_id!=user.id: raise HTTPException(404,'Notification not found')
    row.read=True; db.commit(); return {'ok':True}

@app.get('/notification-preferences')
def preferences(user=Depends(current_user),db=Depends(get_db)):
    row=db.get(NotificationPreference,user.id); return row.data if row else Preferences().model_dump()

@app.put('/notification-preferences')
def preferences_put(body:Preferences,user=Depends(current_user),db=Depends(get_db)):
    row=db.get(NotificationPreference,user.id)
    if row: row.data=body.model_dump()
    else: db.add(NotificationPreference(user_id=user.id,data=body.model_dump()))
    db.commit(); return body

@app.get('/discovery')
def discovery_status(user=Depends(current_user),db=Depends(get_db)):
    return [{'id':r.id,'status':r.status,'started_at':r.started_at,'finished_at':r.finished_at,'details':r.details} for r in db.scalars(select(DiscoveryRun).order_by(DiscoveryRun.started_at.desc()).limit(10))]

@app.post('/discovery')
def request_discovery(user=Depends(current_user)):
    rate_limit('discovery:global',1,300)
    if not(settings.greenhouse_boards or settings.lever_sites): raise HTTPException(503,'Configure GREENHOUSE_BOARDS or LEVER_SITES before discovering postings')
    redis.set('discovery:requested','1',ex=3600)
    return {'status':'queued'}

@app.get('/search-queries')
def search_queries(user=Depends(current_user),db=Depends(get_db)):
    _,p=profile_for(db,user.id)
    # Dynamic deterministic baseline, no invented profile details or open-position claims.
    fields=p.titles or p.skills or ['internship']
    return {'queries':[' '.join(str(v) for v in (field,'internship',p.season,p.graduation.year if p.graduation else None,p.standing_at_start, ' '.join(p.locations or [])) if v) for field in fields[:12]],'executed':False,'note':'Search-engine discovery is not configured. These are query suggestions, not discovered openings.'}

class AnalystQuery(__import__('pydantic').BaseModel):
    query: str = __import__('pydantic').Field(min_length=1,max_length=1000)

@app.post('/analyst')
def analyst(body:AnalystQuery,user=Depends(current_user),db=Depends(get_db)):
    # Extractive retrieval keeps every conclusion tied to a stored original posting.
    words=[w for w in body.query.split() if len(w)>3][:8]
    query=select(Job).where(Job.status=='open')
    from sqlalchemy import or_
    if words: query=query.where(or_(*[Job.text.ilike('%'+w.replace('%','').replace('_','')+'%') for w in words]))
    answers=[]
    for job in db.scalars(query.order_by(Job.retrieved_at.desc()).limit(8)):
        answers.append({'job_id':job.id,'company':job.company,'title':job.title,'url':job.url,'eligibility':evaluate_job(db,user.id,job)})
    db.commit()
    return {'answer':'Source-grounded results from the stored catalogue. Inspect each requirement before applying.' if answers else 'No stored postings match this query. This does not prove there are no openings.','citations':answers,'method':'extractive keyword retrieval; conversational reasoning is not yet implemented'}

@app.get('/account/export')
def export_account(user=Depends(current_user),db=Depends(get_db)):
    _,profile=profile_for(db,user.id)
    db.add(Audit(user_id=user.id,action='account.exported')); db.commit()
    return {'email':user.email,'profile':profile.model_dump(mode='json'),'resumes':[{'id':x.id,'filename':x.filename,'confirmed':x.confirmed,'document':decrypt(x.encrypted_document),'extraction':decrypt(x.encrypted_extraction)} for x in db.scalars(select(Resume).where(Resume.user_id==user.id))],'applications':applications(user,db),'notifications':notifications(user,db),'notification_preferences':preferences(user,db),'saved_jobs':[{'job_id':x.job_id,'hidden':x.hidden} for x in db.scalars(select(SavedJob).where(SavedJob.user_id==user.id))]}

@app.delete('/account')
def delete_account(response:Response,user=Depends(current_user),db=Depends(get_db)):
    # Explicit deletion is portable to SQLite tests as well as PostgreSQL CASCADE.
    for model in (AuthSession,Profile,Qualification,Resume,Evaluation,MatchScore,SavedJob,Application,Notification,NotificationPreference):
        db.execute(delete(model).where(model.user_id==user.id))
    db.execute(delete(Audit).where(Audit.user_id==user.id)); db.delete(db.get(User,user.id)); db.commit()
    response.delete_cookie('internmatch_session'); return {'ok':True}
